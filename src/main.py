# -*- coding: utf-8 -*-
"""资金台账AI自动化 - 主程序
流水线：读取多账户流水 -> 清洗 -> 智能分类入账 -> 日报/月报 -> 异常预警 -> 可视化看板 -> 输出Excel
运行：python main.py  （依赖：pandas / openpyxl / matplotlib）
"""
import pandas as pd
from data_cleaner import load_all_flows
from classifier import classify_frame
from reporter import daily_report, monthly_report
from anomaly import detect
from visualizer import build_charts
from config import OUTPUT_DIR, LEDGER_COLUMNS
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    print("[1/6] 读取并清洗多账户银行流水 ...")
    raw = load_all_flows()
    print(f"      原始流水 {len(raw)} 条，来源账户 {raw['账户'].nunique()} 个")

    print("[2/6] 智能分类入账 ...")
    ledger = classify_frame(raw.copy())
    total = len(ledger)
    unmatched = int((ledger["收支类别"] == "未识别").sum())
    matched = total - unmatched
    print(f"      自动分类 {matched} 条 ({matched / total * 100:.1f}%)，待人工复核 {unmatched} 条")

    print("[3/6] 生成资金日报 / 月报 ...")
    daily = daily_report(ledger)
    monthly = monthly_report(ledger)

    print("[4/6] 异常与风险预警 ...")
    anomalies = detect(ledger)
    hi = int((anomalies["风险等级"] == "高").sum())
    mid = int((anomalies["风险等级"] == "中").sum())
    lo = int((anomalies["风险等级"] == "低").sum())
    print(f"      命中预警 {len(anomalies)} 条（高 {hi} / 中 {mid} / 低 {lo}）")

    print("[5/6] 生成可视化看板 ...")
    exp = ledger[ledger["借贷方向"] == "贷"].groupby("会计科目")["贷方金额"].sum()
    pie_df = exp.reset_index().rename(columns={"会计科目": "科目", "贷方金额": "金额"})\
        .sort_values("金额", ascending=False).head(10)
    charts = build_charts(monthly, pie_df)

    print("[6/6] 输出 Excel 资金台账 ...")
    out = OUTPUT_DIR / "资金台账_2026.xlsx"
    write_excel(out, raw, ledger, daily, monthly, anomalies, charts)
    print(f"完成 -> {out}")


def write_excel(path, raw, ledger, daily, monthly, anomalies, charts):
    raw_w = raw.copy()
    raw_w["记账日期"] = raw_w["记账日期"].dt.strftime("%Y-%m-%d")
    ledger_w = ledger[LEDGER_COLUMNS].copy()
    ledger_w["记账日期"] = ledger_w["记账日期"].dt.strftime("%Y-%m-%d")
    daily_w = daily.copy()
    daily_w["日期"] = daily_w["日期"].astype(str)

    info = pd.DataFrame({"资金台账AI自动化系统": [
        "项目：基于 Python 的多银行账户资金台账自动化系统",
        "流水线：多源流水 -> 清洗 -> 智能分类入账 -> 日报/月报 -> 异常预警 -> 可视化看板",
        "Sheet 导航：",
        "  1) 标准资金台账：清洗 + 自动分类后的统一台账（核心交付）",
        "  2) 原始流水：三家银行原始数据（列名/格式各异，已清洗）",
        "  3) 资金日报 / 4) 资金月报：按日/月聚合的收支汇总",
        "  5) 异常预警：大额 / 借贷不平 / 疑似重复 / 非工作日交易",
        "  6) 可视化看板：收支趋势与支出结构图",
        "说明：本文件由脚本自动生成，所有数据均为脱敏虚构，仅用于演示。",
    ]})

    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        info.to_excel(xw, sheet_name="项目说明", index=False)
        raw_w.to_excel(xw, sheet_name="原始流水", index=False)
        ledger_w.to_excel(xw, sheet_name="标准资金台账", index=False)
        daily_w.to_excel(xw, sheet_name="资金日报", index=False)
        monthly.to_excel(xw, sheet_name="资金月报", index=False)
        anomalies.to_excel(xw, sheet_name="异常预警", index=False)
        pd.DataFrame().to_excel(xw, sheet_name="可视化看板", index=False)

    wb = openpyxl.load_workbook(path)
    style_sheet(wb["项目说明"], info, single_col=True)
    style_sheet(wb["原始流水"], raw_w)
    style_sheet(wb["标准资金台账"], ledger_w)
    style_sheet(wb["资金日报"], daily_w)
    style_sheet(wb["资金月报"], monthly)
    style_sheet(wb["异常预警"], anomalies)

    ws = wb["可视化看板"]
    ws["A1"] = "资金可视化看板（AI 自动生成）"
    ws["A1"].font = Font(bold=True, size=14, color="1F4E78")
    row = 3
    for key in ("trend", "pie"):
        img = XLImage(str(charts[key]))
        ws.add_image(img, f"A{row}")
        row += 24
    wb.save(path)


def style_sheet(ws, df, single_col=False):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="D9D9D9")
    Border(left=thin, right=thin, top=thin, bottom=thin)
    for c in range(1, len(df.columns) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for c in range(1, len(df.columns) + 1):
        col = get_column_letter(c)
        if single_col:
            ws.column_dimensions[col].width = 80
            continue
        maxlen = max([len(str(df.columns[c - 1]))] +
                     [len(str(v)) for v in df.iloc[:, c - 1].head(200)])
        ws.column_dimensions[col].width = min(max(maxlen + 2, 10), 42)
    ws.freeze_panes = "A2"


if __name__ == "__main__":
    main()
